from __future__ import annotations

import json
import re
import zipfile
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Iterable


STANDARD_SECTIONS = {
    "education": ("education",),
    "experience": ("experience", "work experience", "professional experience"),
    "projects": ("projects", "selected projects"),
    "skills": ("technical skills", "skills", "technologies"),
}

ACTION_VERBS = {
    "analyzed", "architected", "automated", "built", "configured", "created",
    "co-created", "combined", "deployed", "designed", "developed", "diagnosed",
    "directed", "documented", "established", "evaluated", "implemented", "improved",
    "integrated", "led", "optimized", "owned", "produced", "reduced",
    "resolved", "retrained", "sourced", "trained", "validated", "verified",
}

STOPWORDS = {
    "a", "an", "and", "are", "as", "at", "be", "by", "for", "from", "has",
    "have", "in", "into", "is", "it", "of", "on", "or", "our", "that", "the",
    "their", "this", "to", "using", "we", "will", "with", "you", "your",
}

DEFAULT_KEYWORDS = [
    "python", "c++", "java", "javascript", "typescript", "go", "rust", "cuda",
    "sql", "postgresql", "mysql", "mongodb", "bigquery", "redis", "kafka",
    "docker", "kubernetes", "linux", "git", "github actions", "ci/cd",
    "aws", "azure", "gcp", "rest api", "microservices", "distributed systems",
    "data structures", "algorithms", "object-oriented programming",
    "automated testing", "unit testing", "integration testing", "regression testing",
    "quality assurance", "debugging", "root cause analysis", "performance optimization",
    "pytorch", "tensorflow", "machine learning", "deep learning", "computer vision",
    "opencv", "tensorrt", "semantic segmentation", "model evaluation", "llm",
    "ros", "ros2", "robotics", "sensor fusion", "slam", "embedded systems",
    "data pipelines", "etl", "data modeling", "ontology", "owl", "shacl", "skos",
    "agile", "scrum", "technical writing", "stakeholder communication",
]


@dataclass
class DocumentInfo:
    path: str
    file_type: str
    text: str
    pages: int | None = None
    tables: int = 0
    text_boxes: int = 0
    images: int = 0
    header_footer_text: bool = False
    minimum_font_points: float | None = None
    parser_notes: list[str] = field(default_factory=list)


@dataclass
class Finding:
    severity: str
    category: str
    message: str
    recommendation: str


@dataclass
class ScanReport:
    resume: str
    parseability_score: int
    format_score: int
    content_score: int
    structural_readiness_score: int
    structural_readiness_max: int
    job_match_score: int | None
    job_match_max: int | None
    overall_score: int | None
    sections_found: list[str]
    contact_fields: dict[str, bool]
    bullet_count: int
    quantified_bullet_count: int
    strong_action_bullet_count: int
    word_count: int
    matched_keywords: list[str]
    missing_keywords: list[str]
    experience_dates_in_order: bool | None
    findings: list[Finding]
    document: DocumentInfo

    def to_dict(self) -> dict:
        return asdict(self)


def _plain_lines(text: str) -> list[str]:
    return [re.sub(r"\s+", " ", line).strip() for line in text.splitlines() if line.strip()]


def extract_pdf(path: Path) -> DocumentInfo:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    page_text = [(page.extract_text() or "") for page in reader.pages]
    images = 0
    for page in reader.pages:
        try:
            images += len(page.images)
        except Exception:
            pass
    notes = []
    if any(len(t.strip()) < 40 for t in page_text):
        notes.append("At least one page produced very little extractable text.")
    return DocumentInfo(
        path=str(path),
        file_type="pdf",
        text="\n".join(page_text),
        pages=len(reader.pages),
        images=images,
        parser_notes=notes,
    )


def extract_docx(path: Path) -> DocumentInfo:
    from docx import Document

    doc = Document(str(path))
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text for cell in row.cells))

    header_footer_text = False
    for section in doc.sections:
        for story in (section.header, section.footer):
            if any(p.text.strip() for p in story.paragraphs):
                header_footer_text = True

    sizes = []
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            size = run.font.size
            if size is None and paragraph.style is not None:
                size = paragraph.style.font.size
            if size is not None:
                sizes.append(round(size.pt, 2))

    text_boxes = 0
    images = 0
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if not name.startswith("word/") or not name.endswith(".xml"):
                continue
            raw = archive.read(name)
            text_boxes += raw.count(b"<w:txbxContent")
            images += raw.count(b"<w:drawing") + raw.count(b"<w:pict")

    return DocumentInfo(
        path=str(path),
        file_type="docx",
        text="\n".join(chunks),
        tables=len(doc.tables),
        text_boxes=text_boxes,
        images=images,
        header_footer_text=header_footer_text,
        minimum_font_points=min(sizes) if sizes else None,
    )


def extract_document(path: Path) -> DocumentInfo:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix in {".txt", ".md"}:
        return DocumentInfo(path=str(path), file_type=suffix[1:], text=path.read_text(encoding="utf-8"))
    raise ValueError(f"Unsupported resume format: {suffix}")


def detect_sections(text: str) -> list[str]:
    lines = {line.lower().strip("# :") for line in _plain_lines(text)}
    found = []
    for canonical, aliases in STANDARD_SECTIONS.items():
        if any(alias in lines for alias in aliases):
            found.append(canonical)
    return found


def detect_contacts(text: str) -> dict[str, bool]:
    lower = text.lower()
    return {
        "email": bool(re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", text, re.I)),
        "phone": bool(re.search(r"(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}", text)),
        "linkedin": "linkedin.com/" in lower,
        "github": "github.com/" in lower,
    }


def extract_bullets(text: str) -> list[str]:
    bullets = []
    for line in _plain_lines(text):
        match = re.match(r"^[•*\-]\s*(.+)", line)
        if match:
            bullets.append(match.group(1).strip())
    return bullets


def _section_slice(text: str, start_name: str, end_names: Iterable[str]) -> str:
    upper = text.upper()
    start = upper.find(start_name.upper())
    if start < 0:
        return ""
    end_positions = [upper.find(name.upper(), start + len(start_name)) for name in end_names]
    end_positions = [pos for pos in end_positions if pos >= 0]
    end = min(end_positions) if end_positions else len(text)
    return text[start:end]


def experience_date_order(text: str) -> tuple[bool | None, list[int]]:
    experience = _section_slice(text, "EXPERIENCE", ("PROJECTS", "TECHNICAL SKILLS", "SKILLS"))
    if not experience:
        return None, []
    endpoints = []
    for line in _plain_lines(experience):
        if not re.search(r"\b(?:19|20)\d{2}\b|\bPresent\b", line, re.I):
            continue
        years = [int(x) for x in re.findall(r"\b(?:19|20)\d{2}\b", line)]
        if re.search(r"\bPresent\b", line, re.I):
            years.append(9999)
        if years:
            endpoints.append(max(years))
    if len(endpoints) < 2:
        return None, endpoints
    return all(a >= b for a, b in zip(endpoints, endpoints[1:])), endpoints


def normalize_keyword(text: str) -> str:
    value = text.lower().replace("–", "-").replace("—", "-")
    value = re.sub(r"\s+", " ", value)
    return value


def keyword_match(resume_text: str, job_text: str, keywords: Iterable[str]) -> tuple[list[str], list[str]]:
    resume = normalize_keyword(resume_text)
    job = normalize_keyword(job_text)
    relevant = [kw for kw in keywords if normalize_keyword(kw) in job]
    matched = [kw for kw in relevant if normalize_keyword(kw) in resume]
    missing = [kw for kw in relevant if normalize_keyword(kw) not in resume]
    return sorted(set(matched)), sorted(set(missing))


def scan_resume(resume_path: Path, job_text: str | None = None, keywords: Iterable[str] | None = None) -> ScanReport:
    document = extract_document(resume_path)
    text = document.text
    lines = _plain_lines(text)
    sections = detect_sections(text)
    contacts = detect_contacts(text)
    bullets = extract_bullets(text)
    quantified = [b for b in bullets if re.search(r"\d", b)]
    strong = [b for b in bullets if b.split() and b.split()[0].lower().rstrip(".,:;") in ACTION_VERBS]
    words = re.findall(r"\b[\w+#./-]+\b", text)
    date_order, endpoints = experience_date_order(text)
    findings: list[Finding] = []

    parseability = 0
    if len(text.strip()) >= 1000:
        parseability += 10
    else:
        findings.append(Finding("high", "parseability", "Too little text was extracted from the file.", "Use a text-based DOCX or PDF rather than an image or scan."))
    if len(lines) >= 20:
        parseability += 5
    if len(sections) == 4:
        parseability += 10
    else:
        missing_sections = sorted(set(STANDARD_SECTIONS) - set(sections))
        findings.append(Finding("high", "structure", f"Standard sections were not all recognized: {', '.join(missing_sections)}.", "Use plain section headings such as Education, Experience, Projects, and Technical Skills."))
    if contacts["email"] and contacts["phone"]:
        parseability += 5
    else:
        findings.append(Finding("high", "contact", "Email or phone number was not recognized.", "Place both in the main document body as plain text."))

    format_score = 25
    if document.tables:
        format_score -= 4
        findings.append(Finding("medium", "format", f"The DOCX contains {document.tables} table(s).", "Use ordinary paragraphs and tab stops where possible; tables can scramble reading order in some parsers."))
    if document.text_boxes:
        format_score -= 6
        findings.append(Finding("high", "format", f"The DOCX contains {document.text_boxes} text box(es).", "Move all resume text into ordinary document paragraphs."))
    if document.header_footer_text:
        format_score -= 4
        findings.append(Finding("high", "format", "Text appears in a header or footer.", "Keep name and contact information in the main document body."))
    if document.minimum_font_points is not None and document.minimum_font_points < 10:
        format_score -= 4
        findings.append(Finding("medium", "readability", f"The smallest detected font is {document.minimum_font_points:g} pt.", "Use at least 10 pt body text; 10.5-11 pt is preferable when space allows."))
    if document.pages is not None and not 1 <= document.pages <= 2:
        format_score -= 4
        findings.append(Finding("medium", "length", f"The PDF has {document.pages} pages.", "For an early-career U.S. industry resume, keep the targeted resume to one page unless the role requires more detail."))
    if date_order is False:
        format_score -= 5
        findings.append(Finding("high", "chronology", f"Experience dates are not in reverse chronological order: {endpoints}.", "Order dated entries from current/most recent to oldest, or split relevant and additional experience into clearly labeled sections."))

    content_score = 0
    if 8 <= len(bullets) <= 24:
        content_score += 5
    else:
        findings.append(Finding("medium", "content", f"Detected {len(bullets)} achievement bullets.", "Use a focused set of relevant accomplishment bullets rather than too few facts or an exhaustive history."))
    action_ratio = len(strong) / len(bullets) if bullets else 0
    if action_ratio >= 0.75:
        content_score += 5
    else:
        findings.append(Finding("medium", "writing", f"Only {action_ratio:.0%} of detected bullets start with a recognized action verb.", "Start bullets with concrete verbs that describe your contribution."))
    metric_ratio = len(quantified) / len(bullets) if bullets else 0
    if metric_ratio >= 0.25:
        content_score += 5
    else:
        findings.append(Finding("low", "evidence", f"Only {metric_ratio:.0%} of bullets contain a number or measurable scope.", "Add truthful scale, performance, coverage, or outcome measures where they clarify the work."))
    if not re.search(r"\b(I|me|my|we|our)\b", text, re.I):
        content_score += 3
    else:
        findings.append(Finding("low", "writing", "First-person pronouns were detected.", "Resume bullets normally omit first-person pronouns."))
    if 350 <= len(words) <= 850:
        content_score += 4
    else:
        findings.append(Finding("medium", "length", f"The resume contains approximately {len(words)} words.", "Keep enough relevant evidence to establish fit without making the page difficult to scan."))
    placeholder_pattern = r"\b(TODO|TBD|lorem ipsum|your name|company 1|employee 1)\b"
    if not re.search(placeholder_pattern, text, re.I):
        content_score += 3
    else:
        findings.append(Finding("high", "content", "Placeholder or synthetic text was detected.", "Replace placeholders with verified information before submission."))

    matched: list[str] = []
    missing: list[str] = []
    job_match_score = None
    overall_score = None
    if job_text:
        matched, missing = keyword_match(text, job_text, keywords or DEFAULT_KEYWORDS)
        relevant_count = len(matched) + len(missing)
        coverage = len(matched) / relevant_count if relevant_count else 1.0
        job_match_score = round(coverage * 20)
        if missing:
            findings.append(Finding("medium", "job match", f"Missing {len(missing)} of {relevant_count} recognized job-specific technical terms.", "Add only supported terms that accurately describe your experience; do not keyword-stuff."))
        overall_score = parseability + max(format_score, 0) + content_score + job_match_score

    return ScanReport(
        resume=str(resume_path),
        parseability_score=parseability,
        format_score=max(format_score, 0),
        content_score=content_score,
        structural_readiness_score=parseability + max(format_score, 0) + content_score,
        structural_readiness_max=80,
        job_match_score=job_match_score,
        job_match_max=20 if job_text else None,
        overall_score=overall_score,
        sections_found=sections,
        contact_fields=contacts,
        bullet_count=len(bullets),
        quantified_bullet_count=len(quantified),
        strong_action_bullet_count=len(strong),
        word_count=len(words),
        matched_keywords=matched,
        missing_keywords=missing,
        experience_dates_in_order=date_order,
        findings=findings,
        document=document,
    )


def report_markdown(report: ScanReport) -> str:
    lines = [
        "# Local ATS Audit",
        "",
        f"Resume: `{report.resume}`",
        "",
        "## Scores",
        "",
        f"- Parseability: {report.parseability_score}/30",
        f"- Format safety: {report.format_score}/25",
        f"- Content evidence: {report.content_score}/25",
        f"- Structural readiness: {report.structural_readiness_score}/{report.structural_readiness_max}",
    ]
    if report.job_match_score is not None:
        lines.extend([
            f"- Job-specific keyword coverage: {report.job_match_score}/20",
            f"- Combined diagnostic score: {report.overall_score}/100",
        ])
    lines.extend([
        "",
        "Scores are local diagnostics, not predictions from an employer's ATS.",
        "",
        "## Extraction summary",
        "",
        f"- Sections recognized: {', '.join(report.sections_found) or 'none'}",
        f"- Contact fields: {json.dumps(report.contact_fields, sort_keys=True)}",
        f"- Words: {report.word_count}",
        f"- Bullets: {report.bullet_count}",
        f"- Quantified bullets: {report.quantified_bullet_count}",
        f"- Strong-action bullets: {report.strong_action_bullet_count}",
        f"- Reverse chronological experience: {report.experience_dates_in_order}",
        "",
        "## Findings",
        "",
    ])
    if not report.findings:
        lines.append("No structural findings.")
    for finding in sorted(report.findings, key=lambda f: {"high": 0, "medium": 1, "low": 2}.get(f.severity, 3)):
        lines.append(f"- **{finding.severity.upper()} - {finding.category}:** {finding.message} {finding.recommendation}")
    if report.job_match_score is not None:
        lines.extend([
            "",
            "## Job keyword evidence",
            "",
            f"- Matched: {', '.join(report.matched_keywords) or 'none'}",
            f"- Missing: {', '.join(report.missing_keywords) or 'none'}",
        ])
    return "\n".join(lines) + "\n"


def save_report(report: ScanReport, json_path: Path | None, markdown_path: Path | None) -> None:
    if json_path:
        json_path.parent.mkdir(parents=True, exist_ok=True)
        json_path.write_text(json.dumps(report.to_dict(), indent=2), encoding="utf-8")
    if markdown_path:
        markdown_path.parent.mkdir(parents=True, exist_ok=True)
        markdown_path.write_text(report_markdown(report), encoding="utf-8")
